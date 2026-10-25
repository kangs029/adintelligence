#!/usr/bin/env python3
import argparse
import json
from collections import Counter

from openpyxl import load_workbook


FIPS_TO_STATE = {
    "01": "AL",
    "02": "AK",
    "04": "AZ",
    "05": "AR",
    "06": "CA",
    "08": "CO",
    "09": "CT",
    "10": "DE",
    "11": "DC",
    "12": "FL",
    "13": "GA",
    "15": "HI",
    "16": "ID",
    "17": "IL",
    "18": "IN",
    "19": "IA",
    "20": "KS",
    "21": "KY",
    "22": "LA",
    "23": "ME",
    "24": "MD",
    "25": "MA",
    "26": "MI",
    "27": "MN",
    "28": "MS",
    "29": "MO",
    "30": "MT",
    "31": "NE",
    "32": "NV",
    "33": "NH",
    "34": "NJ",
    "35": "NM",
    "36": "NY",
    "37": "NC",
    "38": "ND",
    "39": "OH",
    "40": "OK",
    "41": "OR",
    "42": "PA",
    "44": "RI",
    "45": "SC",
    "46": "SD",
    "47": "TN",
    "48": "TX",
    "49": "UT",
    "50": "VT",
    "51": "VA",
    "53": "WA",
    "54": "WV",
    "55": "WI",
    "56": "WY",
    "72": "PR",
}

STATE_NAME_TO_ABBREV = {
    "ALABAMA": "AL",
    "ALASKA": "AK",
    "ARIZONA": "AZ",
    "ARKANSAS": "AR",
    "CALIFORNIA": "CA",
    "COLORADO": "CO",
    "CONNECTICUT": "CT",
    "DELAWARE": "DE",
    "DISTRICT OF COLUMBIA": "DC",
    "FLORIDA": "FL",
    "GEORGIA": "GA",
    "HAWAII": "HI",
    "IDAHO": "ID",
    "ILLINOIS": "IL",
    "INDIANA": "IN",
    "IOWA": "IA",
    "KANSAS": "KS",
    "KENTUCKY": "KY",
    "LOUISIANA": "LA",
    "MAINE": "ME",
    "MARYLAND": "MD",
    "MASSACHUSETTS": "MA",
    "MICHIGAN": "MI",
    "MINNESOTA": "MN",
    "MISSISSIPPI": "MS",
    "MISSOURI": "MO",
    "MONTANA": "MT",
    "NEBRASKA": "NE",
    "NEVADA": "NV",
    "NEW HAMPSHIRE": "NH",
    "NEW JERSEY": "NJ",
    "NEW MEXICO": "NM",
    "NEW YORK": "NY",
    "NORTH CAROLINA": "NC",
    "NORTH DAKOTA": "ND",
    "OHIO": "OH",
    "OKLAHOMA": "OK",
    "OREGON": "OR",
    "PENNSYLVANIA": "PA",
    "RHODE ISLAND": "RI",
    "SOUTH CAROLINA": "SC",
    "SOUTH DAKOTA": "SD",
    "TENNESSEE": "TN",
    "TEXAS": "TX",
    "UTAH": "UT",
    "VERMONT": "VT",
    "VIRGINIA": "VA",
    "WASHINGTON": "WA",
    "WEST VIRGINIA": "WV",
    "WISCONSIN": "WI",
    "WYOMING": "WY",
    "PUERTO RICO": "PR",
}


def clean(value):
    return str(value or "").strip()


def state_to_abbrev(value):
    normalized = clean(value).upper()
    if len(normalized) == 2:
        return normalized
    return STATE_NAME_TO_ABBREV.get(normalized, "")


def row_dicts(path, header_row=1):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = None
    for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if index < header_row:
            continue
        if index == header_row:
            headers = [clean(cell) for cell in row]
            continue
        yield {header: row[position] if position < len(row) else None for position, header in enumerate(headers)}


def load_principal_city_crosswalk(path):
    crosswalk = {}
    for row in row_dicts(path, header_row=3):
        cbsa = clean(row.get("CBSA Code")).zfill(5)
        city = clean(row.get("Principal City Name")).upper()
        state = FIPS_TO_STATE.get(clean(row.get("FIPS State Code")).zfill(2), "")
        if cbsa and city and state:
            crosswalk[f"{city}|{state}"] = f"31000US{cbsa}"
    return crosswalk


def certified_disclosure_cases(path):
    cases = {}
    certified_rows = 0
    for row in row_dicts(path):
        if clean(row.get("VISA_CLASS")) != "H-1B":
            continue
        if not clean(row.get("CASE_STATUS")).startswith("Certified"):
            continue
        case_number = clean(row.get("CASE_NUMBER"))
        employer = clean(row.get("EMPLOYER_NAME"))
        if case_number and employer:
            cases[case_number] = employer
            certified_rows += 1
    return cases, certified_rows

def aggregate(disclosure_path, worksites_path, principal_cities_path):
    crosswalk = load_principal_city_crosswalk(principal_cities_path)
    cases, certified_rows = certified_disclosure_cases(disclosure_path)
    market_counts = {}
    worksite_rows = 0
    certified_worksite_rows = 0
    matched_rows = 0
    unmatched_rows = 0

    for row in row_dicts(worksites_path):
        worksite_rows += 1
        employer = cases.get(clean(row.get("CASE_NUMBER")))
        if not employer:
            continue

        certified_worksite_rows += 1
        city = clean(row.get("WORKSITE_CITY")).upper()
        state = state_to_abbrev(row.get("WORKSITE_STATE"))
        market_geoid = crosswalk.get(f"{city}|{state}")
        if not market_geoid:
            unmatched_rows += 1
            continue

        matched_rows += 1
        market = market_counts.setdefault(market_geoid, {"lcaFilings": 0, "employers": Counter()})
        market["lcaFilings"] += 1
        market["employers"][employer] += 1

    aggregates = []
    for market_geoid, market in sorted(market_counts.items()):
        employers = sorted(market["employers"].items(), key=lambda item: (-item[1], item[0]))[:10]
        aggregates.append(
            {
                "marketGeoid": market_geoid,
                "lcaFilings": market["lcaFilings"],
                "topEmployers": [name for name, _count in employers],
            }
        )

    return {
        "aggregates": aggregates,
        "coverage": {
            "certifiedDisclosureRows": certified_rows,
            "certifiedWorksiteRows": certified_worksite_rows,
            "crosswalkCities": len(crosswalk),
            "matchedRows": matched_rows,
            "unmatchedRows": unmatched_rows,
            "worksiteRows": worksite_rows,
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--disclosure", required=True)
    parser.add_argument("--worksites", required=True)
    parser.add_argument("--principal-cities", required=True)
    args = parser.parse_args()
    print(json.dumps(aggregate(args.disclosure, args.worksites, args.principal_cities), sort_keys=True))


if __name__ == "__main__":
    main()
